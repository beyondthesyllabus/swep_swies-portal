import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

HEADERS = ["Reg. Number", "Full Name", "Sessions Present", "Frequency (%)", "Total Score"]


def _rows(summaries):
    return [
        [s["reg_no"], s["full_name"], f'{s["sessions_present"]}/{s["sessions_total"]}',
         f'{s["frequency_percentage"]}%', s["total_score"]]
        for s in summaries
    ]


def export_register_excel(summaries, title="SWEP/SWIES Final Result"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Final Result"

    ws.merge_cells("A1:E1")
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    fill = PatternFill(start_color="14213D", end_color="14213D", fill_type="solid")
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(_rows(summaries), start=4):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    for col_letter, width in zip("ABCDE", (18, 30, 16, 14, 12)):
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_register_pdf(summaries, title="SWEP/SWIES Final Result"):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    data = [HEADERS] + _rows(summaries)
    table = Table(data, repeatRows=1, colWidths=[3.2 * cm, 6 * cm, 3 * cm, 3 * cm, 2.8 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14213D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (2, 1), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return buffer
